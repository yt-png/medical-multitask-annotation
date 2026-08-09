"""Filesystem and Excel I/O helpers for the annotation pipeline."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ALLOWED_IMAGE_SUFFIXES = frozenset({".jpg", ".jpeg"})

REQUIRED_EXCEL_COLUMNS = ("image_name", "diagnosis_text")


def list_image_files(images_dir: Path | str) -> list[Path]:
    """List allowed image files in a directory as absolute paths.

    Only ``.jpg`` / ``.jpeg`` (case-insensitive) are allowed. Any other
    file in the directory causes a ``ValueError`` (batch failure).
    """

    directory = Path(images_dir)
    if not directory.is_dir():
        raise FileNotFoundError(f"images directory not found: {directory}")

    files = sorted(p for p in directory.iterdir() if p.is_file())
    if not files:
        raise ValueError(f"no files found in images directory: {directory}")

    allowed: list[Path] = []
    for path in files:
        suffix = path.suffix.lower()
        if suffix not in ALLOWED_IMAGE_SUFFIXES:
            raise ValueError(
                f"unsupported file in images directory: {path.name} "
                f"(only .jpg/.jpeg allowed)"
            )
        allowed.append(path.resolve())

    return allowed


def read_diagnosis_excel(excel_path: Path | str) -> list[tuple[str, str]]:
    """Read ``image_name`` / ``diagnosis_text`` rows from the first sheet.

    Empty diagnosis text raises ``ValueError``. Missing required columns
    raise ``ValueError``.
    """

    path = Path(excel_path)
    if not path.is_file():
        raise FileNotFoundError(f"excel file not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError(f"excel sheet is empty: {path}") from exc

        header_cells = [("" if c is None else str(c).strip()) for c in header]
        try:
            name_idx = header_cells.index("image_name")
            text_idx = header_cells.index("diagnosis_text")
        except ValueError as exc:
            raise ValueError(
                "excel must contain columns 'image_name' and 'diagnosis_text'; "
                f"got {header_cells!r}"
            ) from exc

        records: list[tuple[str, str]] = []
        for row_number, row in enumerate(rows, start=2):
            if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
                continue

            raw_name = row[name_idx] if name_idx < len(row) else None
            raw_text = row[text_idx] if text_idx < len(row) else None

            image_name = "" if raw_name is None else str(raw_name).strip()
            if not image_name:
                raise ValueError(f"empty image_name at excel row {row_number}")

            if raw_text is None or str(raw_text).strip() == "":
                raise ValueError(
                    f"empty diagnosis_text for image_name={image_name!r} "
                    f"at excel row {row_number}"
                )

            records.append((image_name, str(raw_text).strip()))

        if not records:
            raise ValueError(f"excel has header but no data rows: {path}")

        return records
    finally:
        workbook.close()


def write_json(path: Path | str, payload: Any) -> None:
    """Write ``payload`` as UTF-8 JSON, creating parent directories as needed."""

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def read_json(path: Path | str) -> Any:
    """Read a UTF-8 JSON file."""

    target = Path(path)
    if not target.is_file():
        raise FileNotFoundError(f"json file not found: {target}")
    return json.loads(target.read_text(encoding="utf-8"))

